from datetime import date
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q, Sum
from django.db.models.functions import TruncMonth
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from accounts.decorators import role_required
from comms.models import Announcement
from events.models import Event, Registration
from finance.models import Transaction
from members.models import Department, Membership
from tasks.models import Task

MANAGERS = ("admin", "board", "head")


def _is_manager(user):
    return user.is_superuser or getattr(user, "role", None) in MANAGERS


def _last_months(n):
    today = timezone.localdate()
    y, m = today.year, today.month
    months = []
    for _ in range(n):
        months.append((y, m))
        m -= 1
        if m == 0:
            m, y = 12, y - 1
    return list(reversed(months))


def _finance_totals():
    approved = Transaction.objects.filter(status=Transaction.Status.APPROVED)
    income = approved.filter(type="income").aggregate(s=Sum("amount"))["s"] or Decimal(0)
    expense = approved.filter(type="expense").aggregate(s=Sum("amount"))["s"] or Decimal(0)
    return income, expense


@login_required
def dashboard(request):
    User = get_user_model()
    now = timezone.now()
    income, expense = _finance_totals()

    stats = {
        "members": User.objects.filter(is_active=True).count(),
        "departments": Department.objects.count(),
        "upcoming_events": Event.objects.filter(end_time__gte=now).count(),
        "balance": income - expense,
        "open_tasks": Task.objects.exclude(status=Task.Status.DONE).count(),
        "overdue_tasks": Task.objects.exclude(status=Task.Status.DONE)
                                     .filter(due_date__lt=timezone.localdate()).count(),
    }

    upcoming = (Event.objects.filter(end_time__gte=now).order_by("start_time")
                .annotate(n_registered=Count("registrations"))[:5])
    my_event_ids = set(Registration.objects.filter(user=request.user)
                       .values_list("event_id", flat=True))
    my_tasks = (Task.objects.filter(assignee=request.user).exclude(status=Task.Status.DONE)
                .select_related("event")[:5])
    announcements = Announcement.objects.select_related("department")[:3]

    # ----- Biểu đồ thu chi 6 tháng -----
    months = _last_months(6)
    first_day = date(months[0][0], months[0][1], 1)
    rows = (Transaction.objects.filter(status=Transaction.Status.APPROVED, date__gte=first_day)
            .annotate(m=TruncMonth("date")).values("m")
            .annotate(income=Sum("amount", filter=Q(type="income")),
                      expense=Sum("amount", filter=Q(type="expense"))))
    by_month = {(r["m"].year, r["m"].month): r for r in rows}
    finance_chart = {
        "labels": [f"{m:02d}/{y}" for y, m in months],
        "income": [float(by_month.get((y, m), {}).get("income") or 0) for y, m in months],
        "expense": [float(by_month.get((y, m), {}).get("expense") or 0) for y, m in months],
    }

    # ----- Biểu đồ đăng ký / điểm danh 6 sự kiện gần nhất -----
    recent = list(Event.objects.filter(start_time__lte=now).order_by("-start_time")
                  .annotate(reg=Count("registrations"),
                            att=Count("registrations",
                                      filter=Q(registrations__checked_in_at__isnull=False)))[:6])
    recent.reverse()
    event_chart = {
        "labels": [e.title[:20] for e in recent],
        "registered": [e.reg for e in recent],
        "attended": [e.att for e in recent],
    }

    # ----- Biểu đồ trạng thái công việc -----
    counts = dict(Task.objects.values_list("status").annotate(n=Count("id")))
    task_chart = {
        "labels": [label for _, label in Task.Status.choices],
        "data": [counts.get(value, 0) for value, _ in Task.Status.choices],
    }

    return render(request, "reports/dashboard.html", {
        "stats": stats,
        "upcoming": upcoming,
        "my_event_ids": my_event_ids,
        "my_tasks": my_tasks,
        "announcements": announcements,
        "finance_chart": finance_chart,
        "event_chart": event_chart,
        "task_chart": task_chart,
        "can_manage": _is_manager(request.user),
    })


# ================= XUẤT EXCEL =================
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="0D6EFD")


def _write_sheet(ws, headers, rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None] + [8])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 50)
    ws.freeze_panes = "A2"


def _fmt_dt(dt):
    return timezone.localtime(dt).strftime("%H:%M %d/%m/%Y") if dt else ""


@role_required(*MANAGERS)
def export_excel(request):
    User = get_user_model()
    wb = Workbook()
    income, expense = _finance_totals()

    # --- Tổng quan ---
    ws = wb.active
    ws.title = "Tổng quan"
    ws.append(["BÁO CÁO HOẠT ĐỘNG CÂU LẠC BỘ"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Ngày xuất: {timezone.localtime():%H:%M %d/%m/%Y} - Người xuất: {request.user}"])
    ws.append([])
    task_counts = dict(Task.objects.values_list("status").annotate(n=Count("id")))
    summary = [
        ("Số thành viên", User.objects.filter(is_active=True).count()),
        ("Số ban", Department.objects.count()),
        ("Tổng số sự kiện", Event.objects.count()),
        ("Tổng lượt đăng ký sự kiện", Registration.objects.count()),
        ("Tổng lượt điểm danh", Registration.objects.filter(checked_in_at__isnull=False).count()),
        ("Tổng thu (đã duyệt)", int(income)),
        ("Tổng chi (đã duyệt)", int(expense)),
        ("Số dư quỹ", int(income - expense)),
    ] + [(f"Công việc - {label}", task_counts.get(value, 0)) for value, label in Task.Status.choices]
    for k, v in summary:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18
    for row in ws.iter_rows(min_row=4, min_col=2, max_col=2):
        for cell in row:
            if isinstance(cell.value, int):
                cell.number_format = "#,##0"

    # --- Thành viên ---
    depts = {}
    for m in Membership.objects.select_related("department"):
        depts.setdefault(m.user_id, []).append(str(m.department))
    users = User.objects.filter(is_active=True).order_by("first_name", "username")
    _write_sheet(wb.create_sheet("Thành viên"),
                 ["Họ tên", "Tài khoản", "MSSV", "Email", "Khoa", "Vai trò", "Ban"],
                 [[str(u), u.username, u.mssv, u.email, u.faculty, u.get_role_display(),
                   ", ".join(depts.get(u.id, []))] for u in users])

    # --- Sự kiện ---
    events = Event.objects.order_by("-start_time").annotate(
        reg=Count("registrations"),
        att=Count("registrations", filter=Q(registrations__checked_in_at__isnull=False)))
    _write_sheet(wb.create_sheet("Sự kiện"),
                 ["Tên sự kiện", "Bắt đầu", "Kết thúc", "Địa điểm", "Ban phụ trách",
                  "Tối đa", "Đăng ký", "Điểm danh", "Tỉ lệ có mặt (%)"],
                 [[e.title, _fmt_dt(e.start_time), _fmt_dt(e.end_time), e.location,
                   str(e.department or ""), e.capacity or "", e.reg, e.att,
                   round(e.att * 100 / e.reg, 1) if e.reg else 0] for e in events])

    # --- Thu chi ---
    trans = Transaction.objects.select_related("category", "created_by", "event")
    ws_f = wb.create_sheet("Thu chi")
    _write_sheet(ws_f,
                 ["Ngày", "Loại", "Danh mục", "Nội dung", "Số tiền (VNĐ)", "Trạng thái",
                  "Sự kiện", "Người tạo"],
                 [[t.date.strftime("%d/%m/%Y"), t.get_type_display(), str(t.category or ""),
                   t.description, int(t.amount), t.get_status_display(), str(t.event or ""),
                   str(t.created_by or "")] for t in trans])
    for row in ws_f.iter_rows(min_row=2, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = "#,##0"

    # --- Công việc ---
    tasks = Task.objects.select_related("assignee", "department", "event")
    _write_sheet(wb.create_sheet("Công việc"),
                 ["Công việc", "Người phụ trách", "Ban", "Sự kiện", "Ưu tiên", "Trạng thái",
                  "Hạn chót", "Quá hạn"],
                 [[t.title, str(t.assignee or ""), str(t.department or ""), str(t.event or ""),
                   t.get_priority_display(), t.get_status_display(),
                   t.due_date.strftime("%d/%m/%Y") if t.due_date else "",
                   "Có" if t.is_overdue else ""] for t in tasks])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    filename = f"ClubHub_BaoCao_{timezone.localdate():%Y%m%d}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
